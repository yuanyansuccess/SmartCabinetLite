# -*- coding: utf-8 -*-
"""代码体检：注释规范 + 字符编码，导出 Excel 检查表。

用法：python tools/code_check.py
"""
import io
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "src")
XLSX = os.path.join(ROOT, "docs", "代码体检报告.xlsx")
EXTS = (".cpp", ".h", ".ui", ".qrc")
SKIP = {"build", "node_modules", ".git", ".vs"}

BANNED = ["已删除", "原来是", "之前是", "已修复", "历史遗留", "改成了",
          "修复：", "重构：", "修改点：", "代码审查修复", "袁总", "领导",
          "TODO", "FIXME"]
MOJIBAKE = ["锟斤拷", "æ", "ä¸", "ï¿½", "�"]
SINGLE_RE = re.compile(r"\b(?:int|qint64|double|float|bool|auto|QString|QStringList)\s+([a-hj-z])\s*[=;]")
SINGLE_ALLOW = set("ijktxyzhgwrb")


def parse_def(line):
    """判定函数定义行；构造函数/析构返回 SKIP，非定义返回 None"""
    if line[0:1] in (" ", "\t"):
        return None
    p = line.rfind("::")
    if p < 0:
        return None
    q = line[p + 2:].find("(")
    if q < 0:
        return None
    name = line[p + 2:p + 2 + q].strip()
    if not re.match(r"^[A-Za-z_]\w*$", name):
        return None
    head = line[:p].split()
    cls = head[-1] if head else ""
    if not cls or not cls[0].isupper() or cls.isupper():
        return None
    if name.startswith("~") or name == cls:
        return "SKIP"
    return name


def has_comment_above(out):
    j = len(out) - 1
    while j >= 0 and not out[j].strip():
        j -= 1
    if j < 0:
        return False
    prev = out[j].strip()
    if prev.endswith("*/"):
        return True
    if not (prev.startswith("//") or prev.startswith("*")):
        return False
    body = prev.lstrip("/").lstrip("*").strip()
    if not body:
        return False
    return not all(ch in "=-_#* \t\u2550\u2500" for ch in body)


files = []
for dirpath, dirs, names in os.walk(SRC):
    dirs[:] = [d for d in dirs if d not in SKIP]
    for fn in sorted(names):
        if fn.endswith(EXTS):
            files.append(os.path.join(dirpath, fn))

rows = []
missing = []
banned_hits = []
enc_hits = []
bom_yes = 0
crlf_src = []
t_total = 0
t_doc = 0

for p in files:
    rel = os.path.relpath(p, ROOT).replace("\\", "/")
    raw = open(p, "rb").read()
    if not raw:
        continue
    if raw[:3] == b"\xef\xbb\xbf":
        bom_yes += 1
    try:
        txt = raw.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        enc_hits.append((rel, 0, "非 UTF-8 文件", str(e)[:50]))
        continue
    lines = txt.splitlines()
    head_ok = "袁燕" in "\n".join(lines[:20])
    if "\r\n" in txt:
        crlf_src.append(rel)
    for w in MOJIBAKE:
        if w in txt:
            enc_hits.append((rel, txt[:txt.index(w)].count("\n") + 1, "疑似乱码字符", w))
            break
    body = txt[1:] if txt.startswith("\ufeff") else txt
    for w in ("\u200b", "\u200c", "\u200d", "\ufeff"):
        if w in body:
            enc_hits.append((rel, body[:body.index(w)].count("\n") + 1, "零宽/隐藏字符", repr(w)))
            break
    if not body.endswith("\n"):
        enc_hits.append((rel, body.count("\n") + 1, "文件末尾缺换行", ""))

    total = 0
    documented = 0
    out = []
    for i, ln in enumerate(lines):
        name = parse_def(ln)
        if name == "SKIP":
            out.append(ln)
            continue
        if name:
            total += 1
            if has_comment_above(out):
                documented += 1
            else:
                missing.append((rel, i + 1, name))
        out.append(ln)
    t_total += total
    t_doc += documented
    for i, ln in enumerate(lines, 1):
        s = ln.strip()
        if s.startswith("//") or s.startswith("*"):
            for w in BANNED:
                if w in s:
                    banned_hits.append((rel, i, s[:60]))
                    break
        if not s.startswith("//"):
            for m in SINGLE_RE.finditer(ln):
                if m.group(1) not in SINGLE_ALLOW:
                    rows.append((rel, i, "无意义单字母变量", m.group(0).strip()))
                    break
    rows.append((rel, 0, "文件头缺作者" if not head_ok else "OK",
                 "%d/%d" % (documented, total)))

print("scanned files: %d" % len(files))
print("funcs: %d documented: %d missing: %d" % (t_total, t_doc, t_total - t_doc))
print("coverage: %.1f%%" % (t_doc * 100.0 / t_total if t_total else 100.0))
print("missing detail: %d  banned: %d  encoding: %d"
      % (len(missing), len(banned_hits), len(enc_hits)))
print("files with BOM: %d   CRLF source files: %d" % (bom_yes, len(crlf_src)))
for it in (missing + banned_hits + enc_hits)[:20]:
    print("   ", it)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "总览"
    ws.append(["文件", "函数注释覆盖", "问题数"])
    for c in range(1, 4):
        ws.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="1F4E79")
    prob = {}
    for rel, _ln, kind, _d in rows:
        if kind != "OK":
            prob[rel] = prob.get(rel, 0) + 1
    for rel, _ln, kind, cover in rows:
        ws.append([rel, cover, prob.get(rel, 0)])
    ws.append([])
    ws.append(["合计", "%d/%d" % (t_doc, t_total),
               len(missing) + len(banned_hits) + len(enc_hits) + len(prob)])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for col, w in zip("ABC", (52, 18, 10)):
        ws.column_dimensions[col].width = w

    for title, data in (("缺失注释", missing), ("违禁表述", banned_hits),
                        ("编码问题", enc_hits)):
        s = wb.create_sheet(title)
        s.append(["文件", "行号", "内容"])
        for c in range(1, 4):
            s.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF")
            s.cell(row=1, column=c).fill = PatternFill("solid", fgColor="C00000")
        for row in data:
            s.append(list(row))
        for col, w in zip("ABC", (52, 8, 62)):
            s.column_dimensions[col].width = w
    wb.save(XLSX)
    print("xlsx ->", XLSX)
except ImportError:
    print("openpyxl not installed, skip xlsx export")
sys.exit(0)
